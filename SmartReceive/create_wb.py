import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo

def create_workbook():
    wb = openpyxl.Workbook()
    sheets = [
        "Dashboard", "Original PL", "First scan", "data Whole",
        "SCAN WHOLE", "PARTIAL", "segregation", "scan partial",
        "data after reception", "OPEN Boxes", "Missing", "Settings"
    ]

    # Remove default sheet
    wb.remove(wb.active)

    for name in sheets:
        wb.create_sheet(name)

    # Setup Settings Table
    ws_settings = wb["Settings"]
    ws_settings.append(["Key", "Value"])
    ws_settings.append(["Organization Name", "SmartWarehouse Corp"])
    ws_settings.append(["Version", "1.0.0"])

    tab = Table(displayName="tblSettings", ref="A1:B3")
    style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                           showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    tab.tableStyleInfo = style
    ws_settings.add_table(tab)

    # Save as .xlsx first
    wb.save("SmartReceive_Pro.xlsx")
    print("Workbook created: SmartReceive_Pro.xlsx")

if __name__ == "__main__":
    create_workbook()
