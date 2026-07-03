import openpyxl
from openpyxl import load_workbook

def validate_workbook(filepath):
    print(f"Validating {filepath}...")
    wb = load_workbook(filepath)
    expected_sheets = [
        "Dashboard", "Original PL", "First scan", "data Whole",
        "SCAN WHOLE", "PARTIAL", "segregation", "scan partial",
        "data after reception", "OPEN Boxes", "Missing", "Settings"
    ]

    missing_sheets = [s for s in expected_sheets if s not in wb.sheetnames]
    if missing_sheets:
        print(f"FAILED: Missing sheets: {missing_sheets}")
    else:
        print("PASSED: All required sheets found.")

    # Check for tables
    all_tables = []
    for sheet in wb.worksheets:
        all_tables.extend(sheet.tables.keys())

    print(f"Found tables: {all_tables}")
    if "tblSettings" in all_tables:
        print("PASSED: tblSettings found.")
    else:
        print("FAILED: tblSettings missing.")

if __name__ == "__main__":
    validate_workbook("SmartReceive_Pro.xlsx")
